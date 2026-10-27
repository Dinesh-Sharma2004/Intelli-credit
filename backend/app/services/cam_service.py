from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas


class CamService:
    def __init__(self, cam_dir: Path):
        self.cam_dir = cam_dir
        self.cam_dir.mkdir(parents=True, exist_ok=True)

    def generate(
        self,
        application: dict,
        documents: list[dict],
        research_hits: list[dict],
        scorecard: dict,
        compliance_checks: list[dict],
        transformer_explainability: dict,
        anomaly_result: dict,
        qualitative_intelligence: dict,
        early_warning_result: dict,
        research_graph: dict,
    ) -> dict:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        safe_name = application["borrower_name"].replace(" ", "_")
        docx_path = self.cam_dir / f"CAM_{safe_name}_{timestamp}.docx"
        pdf_path = self.cam_dir / f"CAM_{safe_name}_{timestamp}.pdf"

        memo_json = self.build_memo_json(
            application,
            documents,
            research_hits,
            scorecard,
            compliance_checks,
            transformer_explainability,
            anomaly_result,
            qualitative_intelligence,
            early_warning_result,
            research_graph,
        )
        self.write_docx(docx_path, memo_json)
        self.write_pdf(pdf_path, memo_json)

        return {
            "docx_path": str(docx_path),
            "pdf_path": str(pdf_path),
            "memo_json": memo_json,
            "generation_status": "completed",
        }

    def build_memo_json(
        self,
        application: dict,
        documents: list[dict],
        research_hits: list[dict],
        scorecard: dict,
        compliance_checks: list[dict],
        transformer_explainability: dict,
        anomaly_result: dict,
        qualitative_intelligence: dict,
        early_warning_result: dict,
        research_graph: dict,
    ) -> dict:
        top_risks = [item for item in scorecard["shap_explanation"] if item["impact"] == "negative_to_credit"][:3]
        top_strengths = [item for item in scorecard["shap_explanation"] if item["impact"] == "positive_to_credit"][:3]
        material_doc_risks = [
            {
                "document_name": doc["document_name"],
                "document_type": doc["document_type"],
                "risk_flags": doc.get("risk_flags", []),
                "confidence_score": doc.get("confidence_score", 0.0),
            }
            for doc in documents
            if doc.get("risk_flags")
        ]
        return {
            "header": {
                "borrower_name": application["borrower_name"],
                "sector": application.get("sector"),
                "requested_amount": application.get("requested_amount"),
                "requested_tenor_months": application.get("requested_tenor_months"),
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "external_id": application.get("external_id"),
                "cin": application.get("cin"),
                "pan": application.get("borrower_pan"),
                "gstin": application.get("borrower_gstin"),
            },
            "executive_summary": {
                "decision": scorecard["decision"],
                "recommended_limit": scorecard["recommended_limit"],
                "recommended_rate": scorecard["recommended_rate"],
                "probability_of_default": scorecard["probability_of_default"],
                "risk_grade": scorecard["risk_grade"],
                "summary": scorecard["explanation_summary"],
                "decision_rationale": self._build_decision_rationale(scorecard),
            },
            "five_cs": scorecard["score_breakdown"],
            "document_summary": [
                {
                    "document_name": doc["document_name"],
                    "document_type": doc["document_type"],
                    "risk_flags": doc.get("risk_flags", []),
                    "confidence_score": doc.get("confidence_score", 0.0),
                }
                for doc in documents
            ],
            "research_summary": [
                {
                    "source_name": hit["source_name"],
                    "title": hit["title"],
                    "summary": hit["summary"],
                    "severity": hit["severity"],
                }
                for hit in research_hits
            ],
            "compliance_summary": compliance_checks,
            "explanations": scorecard["shap_explanation"],
            "transformer_explainability": transformer_explainability,
            "anomaly_summary": anomaly_result,
            "qualitative_intelligence": qualitative_intelligence,
            "early_warning_summary": early_warning_result,
            "research_intelligence_summary": research_graph,
            "scenario_analysis": scorecard.get("scenario_analysis", {}),
            "covenant_recommendations": scorecard.get("covenant_recommendations", []),
            "sanction_simulation": scorecard.get("sanction_simulation", {}),
            "key_observations": {
                "top_risks": top_risks,
                "top_strengths": top_strengths,
                "material_document_risks": material_doc_risks,
                "research_alerts": [hit for hit in research_hits if hit.get("severity") in {"high", "medium"}],
            },
        }

    def write_docx(self, target: Path, memo_json: dict) -> None:
        doc = Document()
        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(10)

        doc.add_heading("Comprehensive Credit Appraisal Memo", level=0)
        header = memo_json["header"]
        metadata_table = doc.add_table(rows=5, cols=2)
        metadata_table.style = "Table Grid"
        metadata_rows = [
            ("Borrower", header["borrower_name"]),
            ("Application ID", header.get("external_id") or "Not assigned"),
            ("Sector", header.get("sector") or "Not provided"),
            ("Requested Facility", self._format_currency(header.get("requested_amount"))),
            ("Requested Tenor", f"{header.get('requested_tenor_months', 0)} months"),
        ]
        for row, values in zip(metadata_table.rows, metadata_rows):
            row.cells[0].text = values[0]
            row.cells[1].text = str(values[1])

        identifiers = [
            ("CIN", header.get("cin")),
            ("PAN", header.get("pan")),
            ("GSTIN", header.get("gstin")),
            ("Generated On", self._format_timestamp(header.get("generated_at"))),
        ]
        doc.add_paragraph("")
        for label, value in identifiers:
            doc.add_paragraph(f"{label}: {value or 'Not available'}")

        summary = memo_json["executive_summary"]
        doc.add_heading("1. Executive Summary", level=1)
        summary_table = doc.add_table(rows=5, cols=2)
        summary_table.style = "Table Grid"
        summary_rows = [
            ("Decision", summary["decision"].upper()),
            ("Risk Grade", summary["risk_grade"]),
            ("Probability of Default", self._format_percent(summary["probability_of_default"])),
            ("Recommended Limit", self._format_currency(summary["recommended_limit"])),
            ("Recommended Rate", self._format_percent(summary["recommended_rate"])),
        ]
        for row, values in zip(summary_table.rows, summary_rows):
            row.cells[0].text = values[0]
            row.cells[1].text = str(values[1])
        doc.add_paragraph(summary["decision_rationale"])
        doc.add_paragraph(summary["summary"])

        doc.add_heading("2. Five Cs Assessment", level=1)
        five_c_table = doc.add_table(rows=1, cols=3)
        five_c_table.style = "Table Grid"
        header_cells = five_c_table.rows[0].cells
        header_cells[0].text = "Dimension"
        header_cells[1].text = "Score"
        header_cells[2].text = "Interpretation"
        for key, value in memo_json["five_cs"].items():
            row = five_c_table.add_row().cells
            row[0].text = key.title()
            row[1].text = f"{value * 10:.2f} / 10"
            row[2].text = self._interpret_five_c_score(value)

        doc.add_heading("3. Document Review", level=1)
        doc.add_paragraph(
            f"Documents reviewed: {len(memo_json['document_summary'])}. "
            f"Material document alerts identified: {len(memo_json['key_observations']['material_document_risks'])}."
        )
        for item in memo_json["document_summary"]:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{item['document_name']} ").bold = True
            paragraph.add_run(
                f"({item['document_type']}, confidence {item['confidence_score']:.2f}) - "
                f"{', '.join(item['risk_flags']) if item['risk_flags'] else 'No explicit document-level risk flags noted.'}"
            )

        doc.add_heading("4. Secondary Research", level=1)
        for item in memo_json["research_summary"]:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{item['source_name']} [{item['severity'].upper()}] ").bold = True
            paragraph.add_run(f"{item['title']}: {item['summary']}")
        doc.add_paragraph(
            f"Promoter risk: {memo_json['research_intelligence_summary'].get('promoter_risk', 0.0):.2f}, "
            f"litigation risk: {memo_json['research_intelligence_summary'].get('litigation_risk', 0.0):.2f}, "
            f"sector outlook: {memo_json['research_intelligence_summary'].get('sector_outlook', 'n/a')}."
        )

        doc.add_heading("5. Regulatory and Compliance Mapping", level=1)
        compliance_table = doc.add_table(rows=1, cols=4)
        compliance_table.style = "Table Grid"
        compliance_headers = compliance_table.rows[0].cells
        compliance_headers[0].text = "Code"
        compliance_headers[1].text = "Regulation"
        compliance_headers[2].text = "Status"
        compliance_headers[3].text = "Assessment"
        for item in memo_json["compliance_summary"]:
            row = compliance_table.add_row().cells
            row[0].text = item["regulation_code"]
            row[1].text = item["regulation_name"]
            row[2].text = item["status"].replace("_", " ").title()
            row[3].text = item["rationale"]

        doc.add_heading("6. Model Explainability and Key Drivers", level=1)
        doc.add_paragraph("XGBoost-style adverse drivers")
        for item in memo_json["key_observations"]["top_risks"]:
            doc.add_paragraph(
                f"{item['feature']}: contribution {item['contribution']:+.4f}",
                style="List Bullet",
            )
        doc.add_paragraph("XGBoost-style positive drivers")
        for item in memo_json["key_observations"]["top_strengths"]:
            doc.add_paragraph(
                f"{item['feature']}: contribution {item['contribution']:+.4f}",
                style="List Bullet",
            )

        doc.add_heading("7. Transformer Evidence from Source Documents", level=1)
        transformer_summary = memo_json["transformer_explainability"]
        doc.add_paragraph(
            f"Transformer evidence method: {transformer_summary.get('method', 'n/a')}. "
            f"Average evidence score: {transformer_summary.get('average_evidence_score', 0.0):.4f}."
        )
        for item in transformer_summary.get("driver_evidence", [])[:4]:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{item['label']} ").bold = True
            paragraph.add_run(f"(top similarity {item.get('top_similarity', 0.0):.4f})")
            for evidence in item.get("evidence", [])[:2]:
                doc.add_paragraph(
                    f"{evidence['document_name']}: {evidence['snippet']}",
                    style="List Bullet",
                )

        doc.add_heading("8. Research-to-Decision Linkage", level=1)
        for item in transformer_summary.get("research_evidence", [])[:4]:
            paragraph = doc.add_paragraph(style="List Bullet")
            source_name = item.get("source_name") or "Research"
            paragraph.add_run(f"{source_name}: {item['label']} ").bold = True
            paragraph.add_run(f"(top similarity {item.get('top_similarity', 0.0):.4f})")
            for evidence in item.get("evidence", [])[:1]:
                doc.add_paragraph(
                    f"{evidence['document_name']}: {evidence['snippet']}",
                    style="List Bullet",
                )

        doc.add_heading("9. Anomaly and Monitoring Summary", level=1)
        anomaly_summary = memo_json["anomaly_summary"]
        doc.add_paragraph(
            f"Anomaly score: {anomaly_summary.get('anomaly_score', 0.0):.2f}. "
            f"Flags: {', '.join(anomaly_summary.get('anomaly_flags', [])) or 'None'}."
        )
        for explanation in anomaly_summary.get("explanations", []):
            doc.add_paragraph(explanation, style="List Bullet")
        warning_summary = memo_json["early_warning_summary"]
        doc.add_paragraph(
            f"Early warning score: {warning_summary.get('early_warning_score', 0.0):.2f}. "
            f"Flags: {', '.join(warning_summary.get('early_warning_flags', [])) or 'None'}."
        )

        doc.add_heading("10. Officer Insight Adjustment Summary", level=1)
        qualitative = memo_json["qualitative_intelligence"]
        doc.add_paragraph(qualitative.get("summary", "No officer insights available."))
        doc.add_paragraph(
            f"Note sentiment: {qualitative.get('sentiment_score', 0.0):.2f}; "
            f"keywords: {', '.join(qualitative.get('keywords_detected', [])) or 'None'}."
        )

        doc.add_heading("11. Scenario Analysis and Covenants", level=1)
        scenario = memo_json["scenario_analysis"]
        if scenario:
            doc.add_paragraph(
                f"Base scenario PD {scenario.get('base', {}).get('pd', 0.0):.2%} with limit {self._format_currency(scenario.get('base', {}).get('limit'))}."
            )
            doc.add_paragraph(
                f"Stress scenario PD {scenario.get('stress', {}).get('pd', 0.0):.2%} with limit {self._format_currency(scenario.get('stress', {}).get('limit'))}."
            )
        for covenant in memo_json.get("covenant_recommendations", []):
            doc.add_paragraph(covenant, style="List Bullet")

        doc.add_heading("12. Credit View", level=1)
        doc.add_paragraph(
            f"The application is assessed as {summary['decision'].upper()} with a recommended sanction limit of "
            f"{self._format_currency(summary['recommended_limit'])} at a pricing indication of "
            f"{self._format_percent(summary['recommended_rate'])}, subject to standard credit, legal, and compliance approvals."
        )

        doc.save(str(target))

    def write_pdf(self, target: Path, memo_json: dict) -> None:
        pdf = canvas.Canvas(str(target), pagesize=A4)
        width, height = A4
        left = 50
        right = width - 50
        y = height - 50

        def write_line(text: str, font: str = "Helvetica", size: int = 10, gap: int = 14) -> None:
            nonlocal y
            pdf.setFont(font, size)
            wrapped = simpleSplit(text, font, size, right - left)
            for line in wrapped:
                if y < 60:
                    pdf.showPage()
                    y = height - 50
                    pdf.setFont(font, size)
                pdf.drawString(left, y, line)
                y -= gap

        write_line("Comprehensive Credit Appraisal Memo", "Helvetica-Bold", 15, 18)
        write_line(f"Borrower: {memo_json['header']['borrower_name']}", "Helvetica-Bold")
        write_line(f"Application ID: {memo_json['header'].get('external_id') or 'Not assigned'}")
        write_line(f"Sector: {memo_json['header'].get('sector') or 'Not provided'}")
        write_line(f"Requested Facility: {self._format_currency(memo_json['header'].get('requested_amount'))}")
        write_line(f"Requested Tenor: {memo_json['header'].get('requested_tenor_months', 0)} months")
        y -= 6

        summary = memo_json["executive_summary"]
        write_line("1. Executive Summary", "Helvetica-Bold", 12, 16)
        write_line(f"Decision: {summary['decision'].upper()}")
        write_line(f"Risk Grade: {summary['risk_grade']}")
        write_line(f"Probability of Default: {self._format_percent(summary['probability_of_default'])}")
        write_line(f"Recommended Limit: {self._format_currency(summary['recommended_limit'])}")
        write_line(f"Recommended Rate: {self._format_percent(summary['recommended_rate'])}")
        write_line(summary["decision_rationale"])
        write_line(summary["summary"])
        y -= 6

        write_line("2. Key Risk Drivers", "Helvetica-Bold", 12, 16)
        for item in memo_json["key_observations"]["top_risks"]:
            write_line(f"- {item['feature']}: contribution {item['contribution']:+.4f}")

        write_line("3. Positive Drivers", "Helvetica-Bold", 12, 16)
        for item in memo_json["key_observations"]["top_strengths"]:
            write_line(f"- {item['feature']}: contribution {item['contribution']:+.4f}")

        write_line("4. Compliance Summary", "Helvetica-Bold", 12, 16)
        for item in memo_json["compliance_summary"]:
            write_line(
                f"- {item['regulation_code']} | {item['status'].replace('_', ' ').title()} | {item['rationale']}"
            )

        write_line("5. Transformer Evidence", "Helvetica-Bold", 12, 16)
        transformer_summary = memo_json["transformer_explainability"]
        write_line(
            f"Method: {transformer_summary.get('method', 'n/a')} | Average evidence score: {transformer_summary.get('average_evidence_score', 0.0):.4f}"
        )
        for item in transformer_summary.get("driver_evidence", [])[:3]:
            write_line(f"- {item['label']} | top similarity {item.get('top_similarity', 0.0):.4f}")
            for evidence in item.get("evidence", [])[:1]:
                write_line(f"  {evidence['document_name']}: {evidence['snippet']}")

        write_line("6. Research Linkage", "Helvetica-Bold", 12, 16)
        for item in transformer_summary.get("research_evidence", [])[:3]:
            write_line(
                f"- {(item.get('source_name') or 'Research')} | {item['label']} | top similarity {item.get('top_similarity', 0.0):.4f}"
            )

        write_line("7. Anomaly and Monitoring", "Helvetica-Bold", 12, 16)
        write_line(
            f"Anomaly score: {memo_json['anomaly_summary'].get('anomaly_score', 0.0):.2f} | "
            f"Early warning score: {memo_json['early_warning_summary'].get('early_warning_score', 0.0):.2f}"
        )
        for flag in memo_json["anomaly_summary"].get("anomaly_flags", [])[:3]:
            write_line(f"- {flag}")

        write_line("8. Covenants", "Helvetica-Bold", 12, 16)
        for covenant in memo_json.get("covenant_recommendations", [])[:3]:
            write_line(f"- {covenant}")

        pdf.save()

    def _format_currency(self, amount: float | int | None) -> str:
        value = float(amount or 0.0)
        return f"INR {value:,.2f}"

    def _format_percent(self, value: float | None) -> str:
        return f"{float(value or 0.0) * 100:.2f}%"

    def _format_timestamp(self, value: str | None) -> str:
        if not value:
            return "Not available"
        return value.replace("T", " ").replace("+00:00", " UTC")

    def _interpret_five_c_score(self, score: float) -> str:
        if score >= 0.75:
            return "Strong"
        if score >= 0.6:
            return "Adequate"
        if score >= 0.45:
            return "Moderate concern"
        return "Weak"

    def _build_decision_rationale(self, scorecard: dict) -> str:
        decision = scorecard["decision"].upper()
        return (
            f"The model recommends {decision} based on the combined assessment of documentary evidence, "
            f"secondary research, qualitative officer inputs, and the explainable score contribution profile."
        )
