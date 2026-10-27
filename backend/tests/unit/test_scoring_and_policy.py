from pathlib import Path

import pytest
from docx import Document

from app.ml.features import FEATURE_KEYS
from app.services.anomaly_service import AnomalyService
from app.services.cam_service import CamService
from app.services.compliance_service import ComplianceService
from app.services.scoring_service import ScoringService

APPLICATION = {"borrower_name": "Acme Foods Private Limited", "sector": "food processing", "requested_amount": 5e7,
               "borrower_pan": "AAPFU0939F", "borrower_gstin": "27AAPFU0939F1ZV"}
DOCUMENT = {
    "document_name": "annual-report.pdf", "document_type": "annual_report", "category": "unstructured_financial",
    "extracted_text": "Revenue Rs. 120 crore, EBITDA margin 14%. Board and audit committee met quarterly. Cash flow healthy.",
    "structured_data": {"currency_amounts_detected": ["120"], "percentages_detected": ["14"], "gstin": "27AAPFU0939F1ZV",
                        "pan": "AAPFU0939F", "cin": None},
    "risk_flags": [], "confidence_score": 0.9,
}


@pytest.mark.parametrize("pd, decision", [(0.01, "approve"), (0.1199, "approve"), (0.12, "review"), (0.2199, "review"), (0.22, "reject"), (0.9, "reject")])
def test_decision_policy_tracks_grade_bands(pd, decision):
    assert ScoringService.decision(pd) == decision


def test_grades_are_monotone_in_pd():
    order = ["AAA", "AA", "A", "BBB+", "BBB", "BB", "B", "CCC", "D"]
    grades = [ScoringService().grade(pd / 100) for pd in range(1, 60)]
    assert [order.index(g) for g in grades] == sorted(order.index(g) for g in grades)


def test_model_scoring_uses_treeshap(pd_model):
    scorecard = ScoringService(pd_model=pd_model).score(APPLICATION, [DOCUMENT], [], {"governance_score": 0.8})
    assert scorecard["explanation_method"] == "xgboost-treeshap"
    assert scorecard["model_version"] == pd_model.version
    drivers = scorecard["shap_explanation"]
    assert [d["feature_key"] for d in drivers if d["feature_key"]] and len(drivers) == len(FEATURE_KEYS)
    assert [d["magnitude"] for d in drivers] == sorted((d["magnitude"] for d in drivers), reverse=True)
    for driver in drivers:
        expected = "negative_to_credit" if driver["shap_value"] > 1e-4 else "positive_to_credit" if driver["shap_value"] < -1e-4 else "neutral"
        assert driver["impact"] == expected
    assert scorecard["decision"] == ScoringService.decision(scorecard["probability_of_default"])
    assert "TreeSHAP" in scorecard["explanation_summary"]


def test_fallback_scorecard_without_model():
    scorecard = ScoringService().score(APPLICATION, [DOCUMENT], [])
    assert scorecard["explanation_method"] == "heuristic-weights"
    assert scorecard["shap_base_value"] is None
    assert 0.02 <= scorecard["probability_of_default"] <= 0.42


def test_feature_vector_is_bounded_and_includes_reconciliation(pd_model):
    anomaly = {"anomaly_score": 0.7, "gst_reconciliation": {"risk_score": 0.65}}
    scorecard = ScoringService(pd_model=pd_model).score(APPLICATION, [DOCUMENT], [], {}, {}, anomaly, {"litigation_risk": 0.9}, {})
    features = scorecard["feature_vector"]
    assert features["gst_reconciliation_risk"] == 0.65
    for key in FEATURE_KEYS:
        assert 0.0 <= features[key] <= 1.0, key
    clean = ScoringService(pd_model=pd_model).score(APPLICATION, [DOCUMENT], [], {}, {}, {"gst_reconciliation": {"risk_score": 0.0}}, {}, {})
    assert clean["feature_vector"]["gst_consistency_score"] > features["gst_consistency_score"]
    assert clean["probability_of_default"] < scorecard["probability_of_default"]


def test_research_litigation_raises_litigation_feature(pd_model):
    service = ScoringService(pd_model=pd_model)
    quiet = service.score(APPLICATION, [DOCUMENT], [], research_graph={"litigation_risk": 0.0})
    noisy = service.score(APPLICATION, [DOCUMENT], [{"finding_type": "litigation", "sentiment": "negative"}],
                          research_graph={"litigation_risk": 0.95})
    assert noisy["feature_vector"]["litigation_risk_score"] > quiet["feature_vector"]["litigation_risk_score"]
    assert noisy["probability_of_default"] >= quiet["probability_of_default"]


def test_compliance_reflects_reconciliation_and_identity():
    names = ["gstr-returns.csv", "bank-statement.csv", "annual-report.pdf"]
    service = ComplianceService()
    clean = {c["regulation_code"]: c for c in service.build_checks(names, [], {"gst_reconciliation": {"flags": []}})}
    assert clean["GST-RECON"]["status"] == "reconciled"
    assert clean["CRA-RATING"]["status"] == "not_applicable"
    flagged = {c["regulation_code"]: c for c in service.build_checks(
        names, [], {"gst_reconciliation": {"flags": ["ITC excess"]}, "identity_checks": {"flags": ["PAN mismatch"]}})}
    assert flagged["GST-RECON"]["status"] == "exception_found" and "ITC excess" in flagged["GST-RECON"]["evidence"]
    assert flagged["RBI-KYC-AML"]["status"] == "exception_found"


def test_identity_checks_detect_pan_and_gstin_problems():
    service = AnomalyService()
    doc = dict(DOCUMENT)
    consistent = service.analyze([doc], APPLICATION)["identity_checks"]
    assert consistent["flags"] == []
    mismatched = service.analyze([doc], APPLICATION | {"borrower_pan": "ZZZPZ9999Z"})["identity_checks"]["flags"]
    assert any("PAN on application not found" in f for f in mismatched)
    assert any("does not embed the borrower PAN" in f for f in mismatched)
    bad_gstin = service.analyze([doc], APPLICATION | {"borrower_gstin": "27AAPFU0939F1ZQ"})["identity_checks"]["flags"]
    assert any("check-digit" in f for f in bad_gstin)


def test_anomaly_prefers_structured_evidence():
    service = AnomalyService()
    narrative = service.analyze([DOCUMENT])
    assert narrative["anomaly_source"] == "narrative_heuristics" and narrative["gst_reconciliation"] is None
    structured_doc = dict(DOCUMENT, structured_data={**DOCUMENT["structured_data"], "gst_returns": [
        {"period": f"2023-{m:02d}", "gstr1_taxable_value": 100, "gstr3b_taxable_value": 70, "itc_available_2b": 10, "itc_claimed_3b": 15}
        for m in range(1, 13)
    ]})
    structured = service.analyze([structured_doc])
    assert structured["anomaly_source"] == "gst_reconciliation"
    assert structured["anomaly_score"] >= structured["gst_reconciliation"]["risk_score"]
    assert any("GSTR-1 vs GSTR-3B" in flag for flag in structured["anomaly_flags"])


def test_cam_generation_writes_docx_and_pdf(tmp_path, pd_model):
    scorecard = ScoringService(pd_model=pd_model).score(APPLICATION, [DOCUMENT], [])
    cam = CamService(tmp_path).generate(APPLICATION | {"external_id": "APP-1"}, [DOCUMENT], [], scorecard, [], {}, {}, {}, {}, {})
    docx_path, pdf_path = Path(cam["docx_path"]), Path(cam["pdf_path"])
    assert docx_path.parent == tmp_path
    assert docx_path.exists() and pdf_path.exists()
    assert pdf_path.read_bytes().startswith(b"%PDF")
    document = Document(str(docx_path))
    text = "\n".join([p.text for p in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert "Acme Foods Private Limited" in text
    assert cam["memo_json"]["executive_summary"]["decision"] == scorecard["decision"]
